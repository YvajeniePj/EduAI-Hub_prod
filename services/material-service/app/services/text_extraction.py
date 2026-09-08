"""
Text extraction service for various file formats
"""
import os
from typing import Optional


def clean_latex_text(raw_text: str) -> str:
    """
    Parse and clean LaTeX text into structured Markdown suitable for educational LLM RAG.
    Preserves math formulas ($...$, $$...$$), structural headers (\section -> ## Раздел: ...),
    definitions, theorems, proofs, and bullet lists.
    """
    import re
    text = raw_text

    # 1. Extract body if \begin{document} is present
    if '\\begin{document}' in text:
        text = text.split('\\begin{document}', 1)[1]
    if '\\end{document}' in text:
        text = text.split('\\end{document}', 1)[0]

    # 2. Remove LaTeX comments (% ...)
    text = re.sub(r'(?m)^%.*$', '', text)

    # 3. Preserve display math environments
    text = re.sub(r'\\begin\{(?:equation|align|gather|multline)\*?\}(.*?)\\end\{(?:equation|align|gather|multline)\*?\}', r'\n$$\1$$\n', text, flags=re.DOTALL)
    text = re.sub(r'\\\[(.*?)\\\]', r'\n$$\1$$\n', text, flags=re.DOTALL)
    text = re.sub(r'\\\((.*?)\\\)', r'$\1$', text, flags=re.DOTALL)

    # 4. Map structural headers into Markdown headers
    text = re.sub(r'\\chapter\*?\{([^}]+)\}', r'\n\n# Глава: \1\n', text)
    text = re.sub(r'\\section\*?\{([^}]+)\}', r'\n\n## Раздел: \1\n', text)
    text = re.sub(r'\\subsection\*?\{([^}]+)\}', r'\n\n### Тема: \1\n', text)
    text = re.sub(r'\\subsubsection\*?\{([^}]+)\}', r'\n\n#### \1\n', text)
    text = re.sub(r'\\paragraph\*?\{([^}]+)\}', r'\n\n**\1:** ', text)

    # 5. Map theorems, definitions, examples
    def _replace_env(match, prefix):
        title = match.group(1)
        if title:
            return f"\n**{prefix} ({title.strip()}):** "
        return f"\n**{prefix}:** "

    text = re.sub(r'\\begin\{(?:definition|defn)\*?\}(?:\[(.*?)\])?', lambda m: _replace_env(m, "Определение"), text)
    text = re.sub(r'\\end\{(?:definition|defn)\*?\}', r'\n', text)
    text = re.sub(r'\\begin\{(?:theorem|thm)\*?\}(?:\[(.*?)\])?', lambda m: _replace_env(m, "Теорема"), text)
    text = re.sub(r'\\end\{(?:theorem|thm)\*?\}', r'\n', text)
    text = re.sub(r'\\begin\{(?:lemma|lem)\*?\}(?:\[(.*?)\])?', lambda m: _replace_env(m, "Лемма"), text)
    text = re.sub(r'\\end\{(?:lemma|lem)\*?\}', r'\n', text)
    text = re.sub(r'\\begin\{(?:example|exmp)\*?\}(?:\[(.*?)\])?', lambda m: _replace_env(m, "Пример"), text)
    text = re.sub(r'\\end\{(?:example|exmp)\*?\}', r'\n', text)
    text = re.sub(r'\\begin\{(?:proof)\*?\}', r'\n*Доказательство:* ', text)
    text = re.sub(r'\\end\{(?:proof)\*?\}', r' ∎\n', text)

    # 6. Map lists and items
    text = re.sub(r'\\begin\{(?:itemize|enumerate|description)\}', r'\n', text)
    text = re.sub(r'\\end\{(?:itemize|enumerate|description)\}', r'\n', text)
    def _replace_item(match):
        label = match.group(1)
        if label:
            return f"\n- **{label.strip()}:** "
        return "\n- "
    text = re.sub(r'\\item(?:\[(.*?)\])?', _replace_item, text)

    # 7. Basic inline styling
    text = re.sub(r'\\textbf\{([^}]+)\}', r'**\1**', text)
    text = re.sub(r'\\textit\{([^}]+)\}', r'*\1*', text)
    text = re.sub(r'\\emph\{([^}]+)\}', r'*\1*', text)
    text = re.sub(r'\\underline\{([^}]+)\}', r'<u>\1</u>', text)
    text = re.sub(r'\\texttt\{([^}]+)\}', r'`\1`', text)

    # 8. Strip non-content commands (labels, refs, citations, styling commands)
    text = re.sub(r'\\(?:cite|ref|eqref|label|input|include|usepackage|documentclass|pagestyle|thispagestyle|geometry)\{[^}]*\}', '', text)
    text = re.sub(r'\\(?:maketitle|tableofcontents|newpage|clearpage|bigskip|medskip|smallskip|noindent)', '', text)

    # 9. Clean residual non-math single backslashes commands while preserving math
    segments = re.split(r'(\$\$.*?\$\$|\$.*?\$)', text, flags=re.DOTALL)
    cleaned_segments = []
    for seg in segments:
        if seg and (seg.startswith('$') or seg.startswith('$$')):
            # This is math, leave it completely untouched!
            cleaned_segments.append(seg)
        elif seg:
            # Clean leftover LaTeX commands in regular prose
            s = re.sub(r'\\[a-zA-Z]+(\[[^\]]*\])?', ' ', seg)
            s = s.replace('{', '').replace('}', '')
            cleaned_segments.append(s)

    text = "".join(cleaned_segments)

    # 10. Normalize whitespace
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n\s*\n\s*\n+', '\n\n', text)
    return text.strip()


def extract_text_from_file(file_path: str, mime_type: str) -> str:
    """Extract text from file of different formats"""
    try:
        if mime_type == "application/pdf" or file_path.lower().endswith('.pdf'):
            try:
                import PyPDF2
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    text = ""
                    for page in pdf_reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            # Try to fix encoding issues for Russian characters
                            try:
                                if 'Ð' in page_text or 'Ñ' in page_text:
                                    page_text = page_text.encode('latin1').decode('utf-8', errors='ignore')
                            except:
                                pass
                            text += page_text + "\n"
                    return text.strip()
            except Exception as pdf_error:
                return f"Error extracting text from PDF: {str(pdf_error)}"
        
        elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or file_path.lower().endswith('.docx'):
            try:
                from docx import Document
                doc = Document(file_path)
                text = ""
                for paragraph in doc.paragraphs:
                    if paragraph.text:
                        text += paragraph.text + "\n"
                return text.strip()
            except Exception as docx_error:
                return f"Error extracting text from DOCX: {str(docx_error)}"
        
        elif mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation" or file_path.lower().endswith('.pptx'):
            try:
                from pptx import Presentation
                prs = Presentation(file_path)
                text = ""
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text:
                            text += shape.text + "\n"
                return text.strip()
            except Exception as pptx_error:
                return f"Error extracting text from PPTX: {str(pptx_error)}"
        
        elif (
            mime_type.startswith("text/")
            or mime_type in ["application/x-tex", "application/x-latex", "application/json", "application/xml", "application/javascript"]
            or any(file_path.lower().endswith(ext) for ext in [
                '.txt', '.tex', '.latex', '.md', '.markdown', '.py', '.js', '.ts', 
                '.html', '.htm', '.css', '.json', '.csv', '.xml', '.yaml', '.yml', 
                '.rtf', '.rst', '.c', '.cpp', '.h', '.java', '.sql', '.sh'
            ])
        ):
            # Try different encodings for text and source files
            encodings = ['utf-8', 'cp1251', 'latin1', 'utf-16']
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        raw = file.read()
                        if file_path.lower().endswith(('.tex', '.latex')) or '\\documentclass' in raw or '\\begin{document}' in raw or '\\section' in raw:
                            return clean_latex_text(raw)
                        return raw.strip()
                except UnicodeDecodeError:
                    continue
            return "Error: could not determine text file encoding"
        
        else:
            # Last fallback: attempt reading as plain text
            encodings = ['utf-8', 'cp1251', 'latin1']
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        content = file.read().strip()
                        if content:
                            if '\\documentclass' in content or '\\begin{document}' in content or '\\section' in content:
                                return clean_latex_text(content)
                            return content
                except Exception:
                    continue
            return f"File format {mime_type} is not supported for text extraction"
    
    except Exception as e:
        return f"Error extracting text: {str(e)}"

